#!/usr/bin/env python3
"""
Builds src/lib/hashtagList.ts (the app's list of allowed #tags) from data/hashtag-taxonomy.xlsx.

Run from proxipoint-mobile/ after editing the spreadsheet:
    python scripts/build_hashtags.py
(needs: pip install openpyxl)

Spreadsheet: sheet "Taxonomy" (or the first sheet), header row with these columns (any order, extra columns ignored):
    Tag ID | Hashtag | Category | Description / Scope | Suggested Icon | Default Radius Sensitivity
- Hashtag: "#LiveMusic" (letters, numbers, _). Unique ignoring case.
- Category: one of CATEGORIES below, or a finer name listed in CATEGORY_ALIASES.
Only Hashtag and Category are required; the other columns can be left out or blank.
- Suggested Icon: an Ionicons-style name ("musical-notes") or a lucide name ("Music"); see ICONS.
  Unknown names fall back to the category's icon (with a warning).
- Default Radius Sensitivity: "Neighborhood (1-3 mi)", "Local (3-5 mi)" or "Metro (5-15 mi)".
"""
import json
import re
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
SOURCE = ROOT / "data" / "hashtag-taxonomy.xlsx"
OUT = ROOT / "src" / "lib" / "hashtagList.ts"

# Category label in the sheet -> id, tint (CategoryBadge colors), default icon.
CATEGORIES = {
    "Music & Nightlife": ("music-nightlife", "pink", "Music"),
    "Food & Dining": ("food-dining", "amber", "Utensils"),
    "Sports & Fitness": ("sports-fitness", "green", "Trophy"),
    "Social & Hobbies": ("social-hobbies", "violet", "Users"),
    "Arts & Culture": ("arts-culture", "violet", "Palette"),
    "Outdoors & Recreation": ("outdoors-recreation", "green", "Trees"),
    "Tech & Professional": ("tech-professional", "blue", "Terminal"),
    "Community & Civic": ("community-civic", "cyan", "HandHeart"),
    "Family & Kids": ("family-kids", "amber", "Users"),
    "Wellness & Mind": ("wellness-mind", "cyan", "Leaf"),
    "Holidays & Seasonal": ("holidays-seasonal", "pink", "Gift"),
    "Celebrations & Milestones": ("celebrations-milestones", "violet", "PartyPopper"),
    "Faith & Spiritual": ("faith-spiritual", "blue", "Church"),
    "Shopping & Markets": ("shopping-markets", "amber", "ShoppingBag"),
}

# Finer category names (as in the "Additional" sheet) -> one of CATEGORIES above. A sheet can use
# either; new names have to be added here (or to CATEGORIES) before they're accepted.
CATEGORY_ALIASES = {
    # Holidays & Seasonal
    "Winter Holidays & Christmas": "Holidays & Seasonal", "Halloween & Fall": "Holidays & Seasonal",
    "New Year's Celebrations": "Holidays & Seasonal", "Thanksgiving & Harvest": "Holidays & Seasonal",
    "Spring & Easter": "Holidays & Seasonal", "Valentine's & Romance": "Holidays & Seasonal",
    "Mardi Gras & Carnival": "Holidays & Seasonal", "St. Patrick's Day": "Holidays & Seasonal",
    "Summer & Patriotic": "Holidays & Seasonal", "Holiday Charity & Giving": "Holidays & Seasonal",
    "Shopping & Holidays": "Holidays & Seasonal", "Seasonal Celebrations": "Holidays & Seasonal",
    "Seasonal & Tours": "Holidays & Seasonal", "Spring & Community": "Holidays & Seasonal",
    # Celebrations & Milestones
    "Weddings & Romance": "Celebrations & Milestones", "Birthdays & Milestones": "Celebrations & Milestones",
    "Family & Milestones": "Celebrations & Milestones", "Education & Milestones": "Celebrations & Milestones",
    "Milestones & Romance": "Celebrations & Milestones", "Career & Milestones": "Celebrations & Milestones",
    "Social & Milestones": "Celebrations & Milestones", "Gatherings & Farewells": "Celebrations & Milestones",
    "Memorial & Remembrance": "Celebrations & Milestones",
    # Faith & Spiritual
    "Family & Faith": "Faith & Spiritual", "Wellness & Spiritual": "Faith & Spiritual",
    # Shopping & Markets
    "Community & Shopping": "Shopping & Markets",
    # Existing categories
    "Sports & Rec": "Sports & Fitness", "Sports & Casual": "Sports & Fitness",
    "Games & Social": "Social & Hobbies", "Makers & Crafts": "Social & Hobbies",
    "Casual Gatherings": "Social & Hobbies", "Home & Hobbies": "Social & Hobbies",
    "Pop Culture & Hobbies": "Social & Hobbies", "Social & Family": "Social & Hobbies",
    "Home & Gatherings": "Social & Hobbies", "Education & Social": "Social & Hobbies",
    "Education & Learning": "Social & Hobbies",
    "Nightlife & Drinks": "Music & Nightlife", "Music & Social": "Music & Nightlife",
    "Food & Casual": "Food & Dining",
    "Cultural Celebrations": "Arts & Culture", "Community & Arts": "Arts & Culture",
    "Arts & Social": "Arts & Culture", "Dance & Social": "Arts & Culture",
    "Outdoors & Tours": "Outdoors & Recreation", "Science & Nature": "Outdoors & Recreation",
    "Professional & Casual": "Tech & Professional", "Professional & Career": "Tech & Professional",
    "Community & Pets": "Community & Civic", "Community & Service": "Community & Civic",
    "Community & Neighborhood": "Community & Civic", "Home & Community": "Community & Civic",
    "Real Estate & Community": "Community & Civic", "Civic & Patriotic": "Community & Civic",
    "Kids & Family": "Family & Kids", "Pets & Family": "Family & Kids",
    "Wellness & Community": "Wellness & Mind",
}

# Sheet icon names -> lucide-react-native components (only ones that exist in lucide 1.x).
# Lucide names written directly in the sheet ("Music") are accepted too if listed in LUCIDE.
ICONS = {
    "musical-notes": "Music", "beer": "Beer", "wine": "Wine", "mic": "Mic", "headset": "Headphones",
    "happy": "PartyPopper", "flash": "Zap", "fast-food": "Truck", "basket": "ShoppingBasket",
    "cafe": "Coffee", "restaurant": "Utensils", "leaf": "Leaf", "bonfire": "Flame",
    "pie-chart": "CakeSlice", "tennisball": "Trophy", "walk": "Footprints", "body": "PersonStanding",
    "bicycle": "Bike", "trail-sign": "Signpost", "baseball": "Trophy", "basketball": "Trophy",
    "football": "Trophy", "fitness": "Mountain", "barbell": "Dumbbell", "dice": "Dices",
    "book": "BookOpen", "color-palette": "Palette", "paw": "PawPrint", "game-controller": "Gamepad2",
    "heart": "Heart", "cut": "Scissors", "chatbubbles": "MessagesSquare", "car": "Car", "film": "Film",
    "camera": "Camera", "color-fill": "PaintBucket", "images": "Images", "library": "Landmark",
    "boat": "Sailboat", "fish": "Fish", "star": "Star", "binoculars": "Binoculars", "snow": "Snowflake",
    "code": "Code", "laptop": "Laptop", "rocket": "Rocket", "business": "Briefcase",
    "hardware-chip": "Cpu", "link": "Link", "briefcase": "Briefcase", "hand-right": "HandHeart",
    "podium": "Landmark", "megaphone": "Megaphone", "trash": "Trash2", "cash": "Banknote",
    "water": "Droplet", "eye": "Eye", "people": "Users", "flask": "FlaskConical", "wind": "Wind",
    "medkit": "Stethoscope", "tent": "Tent", "drama": "Drama",
}
LUCIDE = sorted(set(ICONS.values()) | {c[2] for c in CATEGORIES.values()} | {"Terminal", "Trees"})


def tag_id(tag: str) -> str:
    """'#NYEParty' -> 'nye-party', '#LiveMusic' -> 'live-music' (acronyms stay together)."""
    return re.sub(r"(?<=[a-z0-9])(?=[A-Z])|(?<=[A-Z])(?=[A-Z][a-z])", "-", tag.lstrip("#")).lower()

# A better icon than the sheet's for a few tags.
ICON_OVERRIDES = {"Theater": "Drama", "PuppetShow": "Drama", "Improv": "Drama", "Camping": "Tent"}

RADIUS = {"neighborhood": "neighborhood", "local": "local", "metro": "metro"}


def main() -> None:
    wb = openpyxl.load_workbook(SOURCE, data_only=True)
    ws = wb["Taxonomy"] if "Taxonomy" in wb.sheetnames else wb.worksheets[0]
    header = None
    rows = []
    for raw in ws.iter_rows(values_only=True):
        cells = [str(v).strip() if v is not None else "" for v in raw]
        if header is None:
            if "Hashtag" in cells and "Category" in cells:
                header = {name: i for i, name in enumerate(cells) if name}
            continue
        if not any(cells):
            continue
        rows.append(cells)
    if header is None:
        sys.exit("No header row with 'Hashtag' and 'Category' found.")

    def col(cells, name):
        i = header.get(name)
        return cells[i] if i is not None and i < len(cells) else ""

    seen = set()
    tags = []
    problems = []
    for cells in rows:
        tag = col(cells, "Hashtag").lstrip("#").strip()
        if not tag:
            continue
        if not re.fullmatch(r"[A-Za-z0-9_]+", tag):
            problems.append(f"#{tag}: only letters, numbers and _ are allowed")
            continue
        if tag.lower() in seen:
            problems.append(f"#{tag}: listed twice")
            continue
        seen.add(tag.lower())
        label = CATEGORY_ALIASES.get(col(cells, "Category"), col(cells, "Category"))
        if label not in CATEGORIES:
            problems.append(f"#{tag}: unknown category {label!r} (add it to CATEGORIES)")
            continue
        cat_id, _, cat_icon = CATEGORIES[label]
        sheet_icon = col(cells, "Suggested Icon")
        icon = ICON_OVERRIDES.get(tag) or ICONS.get(sheet_icon) or (sheet_icon if sheet_icon in LUCIDE else None)
        if not icon:
            if sheet_icon:
                print(f"warning: #{tag}: icon {sheet_icon!r} not mapped, using {cat_icon}", file=sys.stderr)
            icon = cat_icon
        radius_raw = col(cells, "Default Radius Sensitivity").split(" ")[0].lower()
        tags.append({
            "id": col(cells, "Tag ID") or tag_id(tag),
            "tag": f"#{tag}",
            "category": cat_id,
            "description": col(cells, "Description / Scope"),
            "icon": icon,
            "radius": RADIUS.get(radius_raw, "local"),
        })
    if problems:
        sys.exit("Fix these rows in the spreadsheet:\n  " + "\n  ".join(problems))

    cats = [
        {"id": cid, "label": label, "tint": tint, "icon": icon}
        for label, (cid, tint, icon) in CATEGORIES.items()
    ]
    icons_used = sorted({t["icon"] for t in tags} | {c["icon"] for c in cats})
    body = [
        "// GENERATED by scripts/build_hashtags.py from data/hashtag-taxonomy.xlsx. Don't edit by hand:",
        "// change the spreadsheet and run `python scripts/build_hashtags.py`.",
        "",
        "export type HashtagCategoryId =",
        *[f"  | '{c['id']}'" for c in cats],
        "  ;",
        "export type HashtagTint = 'cyan' | 'amber' | 'blue' | 'green' | 'violet' | 'pink';",
        "export type HashtagRadius = 'neighborhood' | 'local' | 'metro';",
        f"export type HashtagIcon = {' | '.join(repr(i) for i in icons_used)};",
        "",
        "export interface HashtagCategory {",
        "  id: HashtagCategoryId;",
        "  label: string;",
        "  tint: HashtagTint;",
        "  icon: HashtagIcon;",
        "}",
        "",
        "export interface HashtagEntry {",
        "  /** Stable id from the sheet (\"live-music\"). */",
        "  id: string;",
        "  /** As shown, with # (\"#LiveMusic\"). */",
        "  tag: string;",
        "  category: HashtagCategoryId;",
        "  description: string;",
        "  icon: HashtagIcon;",
        "  /** Default Radius Sensitivity: neighborhood 1-3 mi, local 3-5 mi, metro 5-15 mi. */",
        "  radius: HashtagRadius;",
        "}",
        "",
        f"export const HASHTAG_CATEGORIES: HashtagCategory[] = {json.dumps(cats, indent=2)};",
        "",
        f"export const HASHTAG_LIST: HashtagEntry[] = {json.dumps(tags, indent=2, ensure_ascii=False)};",
        "",
    ]
    text = "\n".join(body)
    # Single quotes, like the rest of the codebase.
    text = re.sub(r'"((?:[^"\\]|\\.)*)"', lambda m: "'" + m.group(1).replace("'", "\\'") + "'", text)
    OUT.write_text(text, encoding="utf-8")
    print(f"Wrote {len(tags)} tags in {len(cats)} categories to {OUT.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
